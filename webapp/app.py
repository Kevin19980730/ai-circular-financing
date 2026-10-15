"""AI Circular Financing -- who funds whom, and who buys from whom, in the AI build-out.

Local Flask app (default http://127.0.0.1:8920).  Everything it shows comes from
files in data/, so it runs with no network and no Claude:

  data/circular_deals.csv      master deal table, one row per leg (money one way,
                               chips/compute the other).  Edit in Excel or via the
                               "Add deal" form.
  data/entities.csv            canonical company names + role (lab, chips, ...)
  data/headline_vs_actual.csv  headline vs disbursed/deployed for the mega deals
  data/candidates.json         news headlines flagged by scripts/scan_news.py
  data/sources/                FT Alphaville / Sona Flourish data, BIS Bulletin 137
                               workbook (refresh: scripts/fetch_sources.py)

    py -3.10 app.py            (PORT env var overrides 8920)

Hosted copy (Render, ACF_HOSTED=1, published by export_site.py into site/webapp/): read-only -
non-GET requests get 403, no news candidates, no Add deal button.
"""
import csv
import datetime as dt
import io
import json
import os
import re
import shutil
import threading
from pathlib import Path

from flask import Flask, Response, jsonify, render_template, request, send_file

HERE = Path(__file__).resolve().parent
DATA = HERE / "data"
SRC = DATA / "sources"
DEALS_CSV = DATA / "circular_deals.csv"
ENTITIES_CSV = DATA / "entities.csv"
HVA_CSV = DATA / "headline_vs_actual.csv"
CANDIDATES_JSON = DATA / "candidates.json"
REVENUE_CSV = DATA / "revenue.csv"
PUBLISHED_JSON = DATA / "published.json"  # written by export_site.py for the hosted copy
HOSTED = os.environ.get("ACF_HOSTED") == "1"

DEAL_FIELDS = [
    "deal_id", "date", "deal", "from", "to", "leg_type", "usd_bn", "leg_note",
    "headline_usd_bn", "headline_note", "bis_category", "status", "confidence",
    "in_totals", "sources", "origin",
]
LEG_TYPES = ["investment", "credit", "backstop", "stock_grant", "purchase", "licensing", "revenue_share"]
BIS_CATEGORIES = [
    "investor_is_counterparty", "supplier_finances_customer",
    "customer_finances_supplier", "reciprocal_commercial",
]
CONFIDENCE = ["high", "medium", "low"]

app = Flask(__name__)
_write_lock = threading.Lock()
_cache = {}


# ---------------------------------------------------------------- loaders --
def _cached(path, loader):
    """Re-read a file only when its mtime changes (edits in Excel show up on reload)."""
    try:
        mtime = path.stat().st_mtime
    except FileNotFoundError:
        return None
    hit = _cache.get(path)
    if hit and hit[0] == mtime:
        return hit[1]
    value = loader(path)
    _cache[path] = (mtime, value)
    return value


def _read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return [dict(r) for r in csv.DictReader(fh)]


def _num(x):
    if x is None:
        return None
    x = str(x).strip().replace(",", "")
    if not x:
        return None
    try:
        return float(x)
    except ValueError:
        return None


def load_deals():
    rows = _cached(DEALS_CSV, _read_csv) or []
    deals, order = {}, []
    for r in rows:
        did = (r.get("deal_id") or "").strip()
        if not did:
            continue
        d = deals.get(did)
        if d is None:
            d = deals[did] = {
                "id": did,
                "date": (r.get("date") or "").strip(),
                "deal": (r.get("deal") or "").strip(),
                "headline_usd_bn": _num(r.get("headline_usd_bn")),
                "headline_note": (r.get("headline_note") or "").strip(),
                "bis_category": (r.get("bis_category") or "").strip(),
                "status": (r.get("status") or "").strip(),
                "confidence": (r.get("confidence") or "").strip(),
                "in_totals": (r.get("in_totals") or "1").strip() not in ("0", "no", "false", "False"),
                "sources": [s.strip() for s in re.split(r"\s*\|\s*", r.get("sources") or "") if s.strip()],
                "origin": (r.get("origin") or "").strip(),
                "legs": [],
            }
            order.append(did)
        if (r.get("from") or "").strip() and (r.get("to") or "").strip():
            d["legs"].append({
                "from": r["from"].strip(),
                "to": r["to"].strip(),
                "type": (r.get("leg_type") or "").strip(),
                "usd_bn": _num(r.get("usd_bn")),
                "note": (r.get("leg_note") or "").strip(),
            })
    return [deals[i] for i in order]


def load_entities():
    rows = _cached(ENTITIES_CSV, _read_csv) or []
    return [
        {"name": r["name"].strip(), "role": (r.get("role") or "other").strip(),
         "layer": (r.get("layer") or "applications").strip(),
         "ticker": (r.get("ticker") or "").strip(), "note": (r.get("note") or "").strip()}
        for r in rows if (r.get("name") or "").strip()
    ]


def load_hva():
    rows = _cached(HVA_CSV, _read_csv) or []
    out = []
    for r in rows:
        out.append({
            "key": r.get("key", ""), "deal": r.get("deal", ""),
            "parties": [s.strip() for s in (r.get("parties") or "").split("|") if s.strip()],
            "announced": r.get("announced", ""),
            "headline_usd_bn": _num(r.get("headline_usd_bn")),
            "actual_usd_bn": _num(r.get("actual_usd_bn")),
            "actual_label": r.get("actual_label", ""), "asof": r.get("asof", ""),
            "status": r.get("status", ""), "note": r.get("note", ""),
            "sources": [s.strip() for s in re.split(r"\s*\|\s*", r.get("sources") or "") if s.strip()],
        })
    return out


def load_revenue():
    """{company: [[fiscal_year_end, usd_bn, basis, source], ...]} sorted by year."""
    out = {}
    for r in _cached(REVENUE_CSV, _read_csv) or []:
        v = _num(r.get("revenue_usd_bn"))
        if v is None:
            continue
        out.setdefault(r["company"].strip(), []).append(
            [r["fiscal_year_end"].strip(), v, r.get("basis", ""), r.get("source", "")])
    for k in out:
        out[k].sort()
    return out


def rev_for(rev, company, date):
    """Revenue of the fiscal year containing `date`; if that year isn't reported yet, the latest one."""
    rows = rev.get(company)
    if not rows:
        return None
    for r in rows:
        if r[0] >= date:
            return r
    return rows[-1]


def _load_json(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def _money_bn(s):
    """'$696.85bn' -> 696.85; 'Undisclosed' -> None."""
    m = re.search(r"\$?([\d.,]+)\s*(bn|tn|m)?", s or "")
    if not m or "ndisclosed" in (s or ""):
        return None
    v = float(m.group(1).replace(",", ""))
    unit = (m.group(2) or "bn").lower()
    return v * 1000 if unit == "tn" else v / 1000 if unit == "m" else v


def load_sona():
    net = _cached(SRC / "sona_radial_network.json", _load_json)
    deals = _cached(SRC / "sona_openai_anthropic_deals.json", _load_json)
    dep = _cached(SRC / "sona_revenue_dependence.json", _load_json)
    out = {"entities": [], "links": 0, "lab_deals": [], "dependence": None}
    if net:
        for n in net["data"]["nodes"]:
            made, recv = (n.get("metadata") or ["", ""])[:2]
            out["entities"].append({"name": n["id"], "made_bn": _money_bn(made), "received_bn": _money_bn(recv)})
        out["links"] = sum(1 for l in net["data"]["links"] if l.get("source"))
        out["fetched"] = net.get("fetched")
    if deals:
        for r in deals["data"]["data"]:
            md = r.get("metadata") or []
            out["lab_deals"].append({
                "lab": r.get("facet"), "from": r.get("source"), "to": r.get("target"),
                "value_label": md[0] if len(md) > 0 else "", "date": md[1] if len(md) > 1 else "",
                "type": md[2] if len(md) > 2 else "", "description": md[3] if len(md) > 3 else "",
                "notes": md[4] if len(md) > 4 else "",
            })
    if dep:
        cols = dep["columns"]["data"]["value"]
        out["dependence"] = {
            "customers": cols,
            "rows": [{"company": r["label"], "values": r["value"]} for r in dep["data"]["data"]],
        }
    return out


def load_bis():
    out = {"split": None, "graph3": None, "graph4": None}
    split = _cached(SRC / "bis_investment_split.json", _load_json)
    if split:
        out["split"] = [
            {"label": r["label"], "pct": r["value"][0]}
            for r in split["data"]["data"] if r["value"] and r["value"][0] is not None
        ]
    xlsx = SRC / "bis_bulletin137_data.xlsx"
    parsed = _cached(xlsx, _parse_bis_xlsx)
    if parsed:
        out.update(parsed)
    return out


def _parse_bis_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    g3, g4 = {}, []
    if "Graph 3" in wb.sheetnames:
        for row in wb["Graph 3"].iter_rows(values_only=True):
            cells = [c for c in row if c is not None]
            # rows look like: label, value, label, value
            for i in range(len(cells) - 1):
                if isinstance(cells[i], str) and isinstance(cells[i + 1], (int, float)):
                    g3[cells[i].strip()] = float(cells[i + 1])
    if "Graph 4" in wb.sheetnames:
        for row in wb["Graph 4"].iter_rows(values_only=True):
            cells = [c for c in row if c is not None]
            if len(cells) == 3 and isinstance(cells[2], (int, float)) and isinstance(cells[0], str):
                g4.append({"source": cells[0], "target": cells[1], "count": int(cells[2])})
    return {"graph3": g3, "graph4": g4}


def load_nortel():
    d = _cached(SRC / "nortel_market_cap.json", _load_json)
    if not d:
        return []
    pts = [(r["label"][:10], r["value"][0]) for r in d["data"]["data"] if r["value"] and r["value"][0] is not None]
    pts.sort()
    return [{"date": a, "bn": round(b / 1e9, 3)} for a, b in pts]


def load_candidates():
    c = _cached(CANDIDATES_JSON, _load_json)
    return c or {"updated": None, "items": []}


# ----------------------------------------------------------------- routes --
@app.before_request
def _read_only():
    if HOSTED and request.method not in ("GET", "HEAD"):
        return jsonify({"ok": False, "errors": ["read-only web copy"]}), 403


@app.route("/")
def index():
    return render_template("index.html", hosted=HOSTED)


@app.route("/api/data")
def api_data():
    deals = load_deals()
    dates = sorted(d["date"] for d in deals if d["date"])
    return jsonify({
        "deals": deals,
        "entities": load_entities(),
        "hva": load_hva(),
        "revenue": load_revenue(),
        "bis": load_bis(),
        "sona": load_sona(),
        "nortel": load_nortel(),
        "candidates": {"updated": None, "items": []} if HOSTED else load_candidates(),
        "meta": {
            "deals_file_updated": dt.datetime.fromtimestamp(DEALS_CSV.stat().st_mtime).isoformat(timespec="minutes")
            if DEALS_CSV.exists() else None,
            "first_date": dates[0] if dates else None,
            "last_date": dates[-1] if dates else None,
            "today": dt.date.today().isoformat(),
            "published": (_cached(PUBLISHED_JSON, _load_json) or {}).get("at") if HOSTED else None,
            "leg_types": LEG_TYPES, "bis_categories": BIS_CATEGORIES,
        },
    })


@app.route("/export.csv")
def export_csv():
    return send_file(DEALS_CSV, mimetype="text/csv", as_attachment=True,
                     download_name="%s_ai_circular_deals.csv" % dt.date.today().strftime("%Y%m%d"))


@app.route("/export.xlsx")
def export_xlsx():
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    deals = load_deals()

    ws = wb.active
    ws.title = "Deals"
    head = ["Date", "Deal", "Headline $bn", "Headline note", "Flows", "BIS category", "Status",
            "Confidence", "In totals", "Sources"]
    ws.append(head)
    for d in sorted(deals, key=lambda x: x["date"], reverse=True):
        flows = "; ".join(
            "%s -> %s %s%s" % (l["from"], l["to"], l["type"],
                               "" if l["usd_bn"] is None else " $%sbn" % _fmt_num(l["usd_bn"]))
            for l in d["legs"])
        ws.append([d["date"], d["deal"], d["headline_usd_bn"], d["headline_note"], flows,
                   d["bis_category"], d["status"], d["confidence"], "yes" if d["in_totals"] else "no",
                   " | ".join(d["sources"])])

    ws2 = wb.create_sheet("Legs")
    ws2.append(["Deal id", "Date", "Deal", "From", "To", "Leg type", "$bn", "Leg note", "BIS category",
                "From: annual revenue $bn", "From: fiscal year end", "% of From revenue",
                "To: annual revenue $bn", "To: fiscal year end", "% of To revenue"])
    rev = load_revenue()
    for d in deals:
        for l in d["legs"]:
            row = [d["id"], d["date"], d["deal"], l["from"], l["to"], l["type"], l["usd_bn"], l["note"], d["bis_category"]]
            for who in (l["from"], l["to"]):
                r = rev_for(rev, who, d["date"])
                pct = round(100 * l["usd_bn"] / r[1], 1) if (r and l["usd_bn"] is not None and r[1]) else None
                row += [r[1] if r else None, r[0] if r else None, pct]
            ws2.append(row)

    ws3 = wb.create_sheet("Headline vs actual")
    ws3.append(["Deal", "Announced", "Headline $bn", "Actual so far $bn", "What actually happened", "As of",
                "Status", "Note", "Sources"])
    for h in load_hva():
        ws3.append([h["deal"], h["announced"], h["headline_usd_bn"], h["actual_usd_bn"], h["actual_label"],
                    h["asof"], h["status"], h["note"], " | ".join(h["sources"])])

    bis = load_bis()
    ws4 = wb.create_sheet("BIS Bulletin 137")
    ws4.append(["BIS Bulletin 137, Circular relationships among AI firms (1 Oct 2026), 2021-25 data"])
    for k, v in (bis.get("graph3") or {}).items():
        ws4.append([k, v])
    ws4.append([])
    ws4.append(["Source layer", "Target layer", "Circular investment count"])
    for r in bis.get("graph4") or []:
        ws4.append([r["source"], r["target"], r["count"]])

    ws6 = wb.create_sheet("Annual revenue")
    ws6.append(["Company", "Fiscal year end", "Revenue $bn", "Basis", "Source"])
    for comp, rows in sorted(rev.items()):
        for r in rows:
            ws6.append([comp] + r)

    sona = load_sona()
    ws5 = wb.create_sheet("FTAV-Sona entities")
    ws5.append(["Entity", "Commitments made $bn", "Commitments received $bn"])
    for e in sorted(sona["entities"], key=lambda e: -((e["made_bn"] or 0) + (e["received_bn"] or 0))):
        ws5.append([e["name"], e["made_bn"], e["received_bn"]])

    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font = Font(bold=True)
        for i, col in enumerate(sheet.columns, 1):
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            sheet.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
        sheet.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     download_name="%s_ai_circular_financing.xlsx" % dt.date.today().strftime("%Y%m%d"))


def _fmt_num(v):
    return ("%.2f" % v).rstrip("0").rstrip(".")


def _slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:40] or "deal"


@app.route("/api/deals", methods=["POST"])
def add_deal():
    """Append one deal (with its legs) to circular_deals.csv."""
    p = request.get_json(force=True, silent=True) or {}
    errors = []
    date = (p.get("date") or "").strip()
    try:
        dt.date.fromisoformat(date)
    except ValueError:
        errors.append("date must be YYYY-MM-DD")
    title = (p.get("deal") or "").strip()
    if not title:
        errors.append("deal title is required")
    cat = (p.get("bis_category") or "").strip()
    if cat not in BIS_CATEGORIES:
        errors.append("pick a BIS category")
    conf = (p.get("confidence") or "medium").strip()
    if conf not in CONFIDENCE:
        errors.append("confidence must be high, medium or low")
    legs = [l for l in (p.get("legs") or []) if (l.get("from") or "").strip() and (l.get("to") or "").strip()]
    if not legs:
        errors.append("add at least one flow (from -> to)")
    for l in legs:
        if l.get("type") not in LEG_TYPES:
            errors.append("flow type must be one of " + ", ".join(LEG_TYPES))
            break
        if (l.get("usd_bn") not in (None, "")) and _num(l.get("usd_bn")) is None:
            errors.append("flow $bn must be a number (or blank if undisclosed)")
            break
    if (p.get("headline_usd_bn") not in (None, "")) and _num(p.get("headline_usd_bn")) is None:
        errors.append("headline $bn must be a number (or blank)")
    if errors:
        return jsonify({"ok": False, "errors": errors}), 400

    existing = {d["id"] for d in load_deals()}
    base = "%s-%s" % (_slug(title), date[:7])
    did, n = base, 2
    while did in existing:
        did, n = "%s-%d" % (base, n), n + 1
    sources = " | ".join(s.strip() for s in re.split(r"[\s|]+", p.get("sources") or "") if s.strip())
    rows = []
    for l in legs:
        rows.append({
            "deal_id": did, "date": date, "deal": title,
            "from": l["from"].strip(), "to": l["to"].strip(), "leg_type": l["type"],
            "usd_bn": "" if _num(l.get("usd_bn")) is None else _fmt_num(_num(l.get("usd_bn"))),
            "leg_note": (l.get("note") or "").strip(),
            "headline_usd_bn": "" if _num(p.get("headline_usd_bn")) is None else _fmt_num(_num(p.get("headline_usd_bn"))),
            "headline_note": (p.get("headline_note") or "").strip(),
            "bis_category": cat, "status": (p.get("status") or "announced").strip(),
            "confidence": conf, "in_totals": "1", "sources": sources, "origin": "manual",
        })
    with _write_lock:
        try:
            shutil.copyfile(DEALS_CSV, DATA / "circular_deals.bak.csv")
            with open(DEALS_CSV, "a", newline="", encoding="utf-8") as fh:
                w = csv.DictWriter(fh, fieldnames=DEAL_FIELDS)
                for r in rows:
                    w.writerow(r)
        except PermissionError:
            return jsonify({"ok": False, "errors": ["circular_deals.csv is locked - close it in Excel and retry"]}), 409
    _mark_candidate(p.get("candidate_id"), "added")
    return jsonify({"ok": True, "deal_id": did})


def _mark_candidate(cid, state):
    if not cid:
        return False
    with _write_lock:
        c = load_candidates()
        hit = False
        for it in c.get("items", []):
            if it.get("id") == cid:
                it["state"] = state
                hit = True
        if hit:
            CANDIDATES_JSON.write_text(json.dumps(c, ensure_ascii=False, indent=1), encoding="utf-8")
    return hit


@app.route("/api/candidates/<cid>/dismiss", methods=["POST"])
def dismiss_candidate(cid):
    return jsonify({"ok": _mark_candidate(cid, "dismissed")})


@app.route("/healthz")
def healthz():
    return "ok"


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "8920"))
    app.run(host="127.0.0.1", port=port, debug=False, threaded=True)
